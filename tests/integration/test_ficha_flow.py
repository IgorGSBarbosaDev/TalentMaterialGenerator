from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app.core.generator_ficha import generate_ficha_pptx
from app.core.reader import lookup_ficha_employees, normalize_filename, read_spreadsheet


def _build_standardized_spreadsheet(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(
        [
            "matricula",
            "nome",
            "idade",
            "cargo",
            "antiguidade",
            "formacao",
            "resumo_perfil",
            "trajetoria",
            "nota_2025",
            "nota_2024",
            "nota_2023",
        ]
    )
    sheet.append(
        [
            "101",
            "Ana Martins",
            "31",
            "Engenheira de Processos",
            "5 anos",
            "Engenharia Metalurgica",
            "Profissional colaborativa e analitica",
            "Analista Jr; Analista Pleno; Especialista",
            "4 / PROM",
            "5 / AP",
            "3 / MN+",
        ]
    )
    sheet.append(
        [
            "102",
            "Carlos Souza",
            "39",
            "Coordenador de Manutencao",
            "8 anos",
            "Engenharia Mecanica",
            "Lidera times multidisciplinares",
            "Tecnico; Supervisor; Coordenador",
            "5 / AP",
            "4 / PROM",
            "3 / MN+",
        ]
    )
    workbook.save(path)
    return path


def _build_new_evaluation_spreadsheet(
    path: Path,
    *,
    with_consolidated_values: bool,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(
        [
            "Matricula",
            "Nome",
            "Cargo",
            "Idade",
            "Antiguidade",
            "Formacao",
            "Resumo do perfil",
            "Trajetoria",
            "Avaliação 2025",
            "Avaliação 2024",
            "Avaliação 2023",
            "Nota 2025",
            "Potencial 2025",
            "Nota 2024",
            "Potencial 2024",
            "Nota 2023",
            "Potencial 2023",
            "CEO1",
            "CEO2",
            "CEO3",
            "CEO4",
        ]
    )
    sheet.append(
        [
            "101",
            "Ana Martins",
            "Engenheira de Processos",
            "31",
            "5 anos",
            "Engenharia Metalurgica",
            "Profissional colaborativa e analitica",
            "Analista Jr; Analista Pleno; Especialista",
            "4 / PROM" if with_consolidated_values else "",
            "5 / AP" if with_consolidated_values else "",
            "",
            "4",
            "PROM",
            "5",
            "AP",
            "",
            "",
            "Diretoria",
            "VP",
            "Sucessor Imediato",
            "Em desenvolvimento",
        ]
    )
    workbook.save(path)
    return path


def test_full_ficha_lookup_and_generation_creates_single_pptx(tmp_path: Path) -> None:
    spreadsheet = _build_standardized_spreadsheet(tmp_path / "ficha.xlsx")
    rows = read_spreadsheet(str(spreadsheet))
    matches = lookup_ficha_employees(rows, name_query="ana")

    generated_file = generate_ficha_pptx(matches[0], str(tmp_path))

    assert Path(generated_file).exists()


def test_generated_slide_has_wide_dimensions(tmp_path: Path) -> None:
    spreadsheet = _build_standardized_spreadsheet(tmp_path / "ficha.xlsx")
    rows = read_spreadsheet(str(spreadsheet))
    matches = lookup_ficha_employees(rows, matricula_query="102")

    generated_file = generate_ficha_pptx(matches[0], str(tmp_path))
    presentation = Presentation(generated_file)

    assert presentation.slide_width == Inches(13.333)
    assert presentation.slide_height == Inches(7.5)


def test_new_evaluation_layout_uses_consolidated_values_when_present(
    tmp_path: Path,
) -> None:
    spreadsheet = _build_new_evaluation_spreadsheet(
        tmp_path / "ficha_new_consolidated.xlsx",
        with_consolidated_values=True,
    )
    rows = read_spreadsheet(str(spreadsheet))

    matches = lookup_ficha_employees(rows, matricula_query="101")

    assert matches[0]["nota_2025"] == "4 / PROM"
    assert matches[0]["nota_2024"] == "5 / AP"
    generated_file = generate_ficha_pptx(matches[0], str(tmp_path))

    assert Path(generated_file).exists()


def test_new_evaluation_layout_falls_back_to_score_and_potential(
    tmp_path: Path,
) -> None:
    spreadsheet = _build_new_evaluation_spreadsheet(
        tmp_path / "ficha_new_fallback.xlsx",
        with_consolidated_values=False,
    )
    rows = read_spreadsheet(str(spreadsheet))

    matches = lookup_ficha_employees(rows, matricula_query="101")

    assert matches[0]["nota_2025"] == "4 / PROM"
    assert matches[0]["nota_2024"] == "5 / AP"
    generated_file = generate_ficha_pptx(matches[0], str(tmp_path))

    assert Path(generated_file).exists()


def test_output_filename_is_normalized_from_selected_employee(tmp_path: Path) -> None:
    employee = {
        "matricula": "77",
        "nome": "Joao Barbara",
        "idade": "",
        "cargo": "Analista",
        "antiguidade": "",
        "formacao": "",
        "resumo_perfil": "",
        "trajetoria": "",
    }

    generated_file = generate_ficha_pptx(employee, str(tmp_path))

    assert Path(generated_file).stem == normalize_filename(employee["nome"])
